from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


WORKSPACE_ROOT = Path(__file__).resolve().parent
OUTPUT_FILE = WORKSPACE_ROOT / "result_check_IGreset_timing.xlsx"

TARGETS = {
    "RADAR": WORKSPACE_ROOT / "RADAR" / "Diag_RADAR_logs",
    "MPC": WORKSPACE_ROOT / "MPC" / "Diag_MPC_logs",
}

TIMESTAMP_PATTERN = re.compile(r"^\s*([0-9]+(?:\.[0-9]+)?)\b")
VSA_2FF_PATTERN = re.compile(r"\bVSA_2FF\b")
DELTA_THRESHOLD = 0.15  # 150 ms

SYSTEM_DID_RULES = {
    "RADAR": {
        "request_id": "18DAB0F1",
        "did_commands": (
            "03 22 A8 C4 00 00 00 00",
            "03 22 A8 C5 00 00 00 00",
        ),
    },
    "MPC": {
        "request_id": "18DAB5F1",
        "did_commands": (
            "03 22 A8 C6 00 00 00 00",
        ),
    },
}


def natural_sort_key(text: str) -> tuple:
    """Sort text by numeric chunks for natural ordering."""
    parts = re.split(r"(\d+)", text)
    return tuple(int(part) if part.isdigit() else part.lower() for part in parts)


def extract_vsa_timestamps(log_path: Path) -> list[float]:
    """Extract all timestamps for VSA_2FF messages in order."""
    timestamps: list[float] = []

    with log_path.open("r", encoding="utf-8", errors="ignore") as file:
        for line in file:
            stripped = line.strip()
            if not stripped:
                continue

            # Check if line contains VSA_2FF
            if not VSA_2FF_PATTERN.search(stripped):
                continue

            # Extract timestamp
            match = TIMESTAMP_PATTERN.match(stripped)
            if match:
                timestamps.append(float(match.group(1)))

    return timestamps


def find_large_deltas(timestamps: list[float]) -> list[tuple[float, float]]:
    """Find consecutive VSA_2FF pairs with delta > 150ms."""
    events: list[tuple[float, float]] = []

    for i in range(len(timestamps) - 1):
        delta = timestamps[i + 1] - timestamps[i]
        if delta > DELTA_THRESHOLD:
            events.append((timestamps[i], timestamps[i + 1]))

    return events


def build_replay_delays_ms(events: list[tuple[float, float]]) -> tuple[list[float], list[float]]:
    """Build CAPL replay waits in milliseconds for each IG reset event.

    - off_delays_ms[i] is wait until IG_OFF_i from previous IG_ON (or from log start for i=0)
    - on_delays_ms[i] is wait from IG_OFF_i to IG_ON_i
    """
    if not events:
        return [], []

    off_delays_ms: list[float] = []
    on_delays_ms: list[float] = []

    for idx, (off_time, on_time) in enumerate(events):
        if idx == 0:
            off_wait = off_time
        else:
            prev_on_time = events[idx - 1][1]
            off_wait = off_time - prev_on_time

        on_wait = on_time - off_time

        off_delays_ms.append(off_wait * 1000.0)
        on_delays_ms.append(on_wait * 1000.0)

    return off_delays_ms, on_delays_ms


def normalize_log_name(system_name: str, log_path: Path) -> str:
    stem = log_path.stem
    for prefix in (
        f"Diag_customer_{system_name}_",
        "Diag_customer_",
    ):
        if stem.startswith(prefix):
            return stem[len(prefix) :]
    return stem


def collect_ig_reset_events(
    system_name: str, folder: Path
) -> dict[str, dict[str, object]]:
    """Collect IG reset events (delta > 150ms) for all logs."""
    results: dict[str, dict[str, object]] = {}

    if not folder.exists():
        return results

    for log_path in sorted(folder.glob("Diag_customer_*.asc")):
        timestamps = extract_vsa_timestamps(log_path)
        events = find_large_deltas(timestamps)
        missing_did_flags = detect_missing_did_after_ig_reset(system_name, log_path, events)
        off_delays_ms, on_delays_ms = build_replay_delays_ms(events)

        log_name = normalize_log_name(system_name, log_path)
        results[log_name] = {
            "events": events,
            "missing_did_flags": missing_did_flags,
            "off_delays_ms": off_delays_ms,
            "on_delays_ms": on_delays_ms,
        }

    return results


def detect_missing_did_after_ig_reset(
    system_name: str,
    log_path: Path,
    events: list[tuple[float, float]],
) -> list[bool]:
    """Return True when a reset window has no required DID request.

    For reset i, the checked window is:
    - from IG_ON_i to IG_OFF_(i+1) for non-last resets
    - from IG_ON_last to end-of-log for the last reset
    """
    if not events:
        return []

    rule = SYSTEM_DID_RULES.get(system_name)
    if rule is None:
        return [False] * len(events)

    request_id = str(rule["request_id"])
    did_commands = tuple(str(cmd) for cmd in rule["did_commands"])
    on_times = [on_time for _, on_time in events]
    next_off_times = [events[idx + 1][0] for idx in range(len(events) - 1)] + [None]
    missing_did_flags = [True] * len(events)
    window_idx = 0

    with log_path.open("r", encoding="utf-8", errors="ignore") as file:
        for line in file:
            if request_id not in line:
                continue

            match = TIMESTAMP_PATTERN.match(line)
            if not match:
                continue

            timestamp = float(match.group(1))
            if not any(command in line for command in did_commands):
                continue

            while window_idx < len(events):
                next_off_time = next_off_times[window_idx]
                if next_off_time is None or timestamp < next_off_time:
                    break
                window_idx += 1

            if window_idx >= len(events):
                continue

            on_time = on_times[window_idx]
            next_off_time = next_off_times[window_idx]
            if timestamp < on_time:
                continue

            if next_off_time is None or timestamp < next_off_time:
                missing_did_flags[window_idx] = False

    return missing_did_flags


def write_sheet(
    workbook: Workbook,
    sheet_name: str,
    logs_data: dict[str, dict[str, object]],
) -> None:
    ws = workbook.create_sheet(title=sheet_name)
    bold_font = Font(bold=True)
    left_alignment = Alignment(horizontal="left")
    delta_fill = PatternFill(fill_type="solid", start_color="C6EFCE", end_color="C6EFCE")
    alert_fill = PatternFill(fill_type="solid", start_color="FFF2CC", end_color="FFF2CC")
    alert_delta_cells = []

    # Find max number of IG reset events across all logs
    max_events = max(
        (len(log_entry["events"]) for log_entry in logs_data.values()),
        default=0,
    )

    # Write headers
    ws["A1"] = "No."
    ws["A1"].font = bold_font
    ws["A1"].alignment = left_alignment

    ws["B1"] = "Log name"
    ws["B1"].font = bold_font
    ws["B1"].alignment = left_alignment

    for event_idx in range(1, max_events + 1):
        off_col = 3 + (event_idx - 1) * 3
        on_col = 4 + (event_idx - 1) * 3
        delta_col = 5 + (event_idx - 1) * 3

        off_cell = ws.cell(row=1, column=off_col)
        on_cell = ws.cell(row=1, column=on_col)
        delta_cell = ws.cell(row=1, column=delta_col)

        off_cell.value = f"IG_OFF_{event_idx}"
        on_cell.value = f"IG_ON_{event_idx}"
        delta_cell.value = f"Delta_{event_idx} (s)"

        off_cell.font = bold_font
        on_cell.font = bold_font
        delta_cell.font = bold_font
        off_cell.alignment = left_alignment
        on_cell.alignment = left_alignment
        delta_cell.alignment = left_alignment

    # Write data rows
    row_num = 2
    for idx, log_name in enumerate(sorted(logs_data.keys(), key=natural_sort_key), start=1):
        log_entry = logs_data[log_name]
        events = log_entry["events"]
        missing_did_flags = log_entry["missing_did_flags"]

        ws.cell(row=row_num, column=1, value=idx)
        ws.cell(row=row_num, column=1).alignment = left_alignment

        ws.cell(row=row_num, column=2, value=log_name)
        ws.cell(row=row_num, column=2).alignment = left_alignment

        for event_idx, (off_time, on_time) in enumerate(events):
            off_col = 3 + event_idx * 3
            on_col = 4 + event_idx * 3
            delta_col = 5 + event_idx * 3

            off_cell = ws.cell(row=row_num, column=off_col)
            on_cell = ws.cell(row=row_num, column=on_col)
            delta_cell = ws.cell(row=row_num, column=delta_col)

            off_cell.value = off_time
            on_cell.value = on_time
            off_ref = ws.cell(row=row_num, column=off_col).coordinate
            on_ref = ws.cell(row=row_num, column=on_col).coordinate
            delta_cell.value = f"={on_ref}-{off_ref}"

            off_cell.alignment = left_alignment
            on_cell.alignment = left_alignment
            delta_cell.alignment = left_alignment

            if event_idx < len(missing_did_flags) and missing_did_flags[event_idx]:
                alert_delta_cells.append(delta_cell)

        row_num += 1

    # Fill all Delta columns from header to last row with data
    last_data_row = max(1, row_num - 1)
    for event_idx in range(1, max_events + 1):
        delta_col = 5 + (event_idx - 1) * 3
        for row_idx in range(1, last_data_row + 1):
            ws.cell(row=row_idx, column=delta_col).fill = delta_fill

    for delta_cell in alert_delta_cells:
        delta_cell.fill = alert_fill

    # Auto-fit column widths
    max_col = 2 + max_events * 3
    ws.auto_filter.ref = f"A1:{get_column_letter(max_col)}1"
    for col_idx in range(1, max_col + 1):
        max_length = 0
        for row in ws.iter_rows(min_row=1, max_row=row_num - 1, min_col=col_idx, max_col=col_idx):
            cell = row[0]
            value = "" if cell.value is None else str(cell.value)
            if len(value) > max_length:
                max_length = len(value)
        ws.column_dimensions[get_column_letter(col_idx)].width = max_length + 2


def write_replay_helper_sheet(
    workbook: Workbook,
    sheet_name: str,
    logs_data: dict[str, dict[str, object]],
) -> None:
    ws = workbook.create_sheet(title=sheet_name)
    bold_font = Font(bold=True)
    left_alignment = Alignment(horizontal="left")

    headers = [
        "No.",
        "Log name",
        "Event",
        "Wait_off_ms",
        "Wait_on_ms",
        "CAPL setTimer(Wait_off, ...) line",
        "CAPL setTimer(Wait_on, ...) line",
    ]

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = bold_font
        cell.alignment = left_alignment

    row_num = 2
    for idx, log_name in enumerate(sorted(logs_data.keys(), key=natural_sort_key), start=1):
        log_entry = logs_data[log_name]
        off_delays_ms = list(log_entry["off_delays_ms"])
        on_delays_ms = list(log_entry["on_delays_ms"])

        if not off_delays_ms:
            continue

        for event_idx, (off_delay_ms, on_delay_ms) in enumerate(
            zip(off_delays_ms, on_delays_ms),
            start=1,
        ):
            ws.cell(row=row_num, column=1, value=idx).alignment = left_alignment
            ws.cell(row=row_num, column=2, value=log_name).alignment = left_alignment
            ws.cell(row=row_num, column=3, value=event_idx).alignment = left_alignment
            ws.cell(row=row_num, column=4, value=off_delay_ms).alignment = left_alignment
            ws.cell(row=row_num, column=5, value=on_delay_ms).alignment = left_alignment
            ws.cell(
                row=row_num,
                column=6,
                value=f"setTimer(Wait_off, {off_delay_ms:.3f});",
            ).alignment = left_alignment
            ws.cell(
                row=row_num,
                column=7,
                value=f"setTimer(Wait_on, {on_delay_ms:.3f});",
            ).alignment = left_alignment
            row_num += 1

    last_data_row = max(1, row_num - 1)
    ws.auto_filter.ref = f"A1:G{last_data_row}"
    for col_idx in range(1, 8):
        max_length = 0
        for row in ws.iter_rows(min_row=1, max_row=last_data_row, min_col=col_idx, max_col=col_idx):
            cell = row[0]
            value = "" if cell.value is None else str(cell.value)
            if len(value) > max_length:
                max_length = len(value)
        ws.column_dimensions[get_column_letter(col_idx)].width = max_length + 2


def main() -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)

    for system_name, folder in TARGETS.items():
        logs_data = collect_ig_reset_events(system_name, folder)
        write_sheet(workbook, system_name, logs_data)
        write_replay_helper_sheet(workbook, f"{system_name}_CAPL", logs_data)

    workbook.save(OUTPUT_FILE)
    print(f"Created: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
